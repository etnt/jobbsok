"""Skapar CV och personligt brev för en annons, redigerbara och exporterbara som Word.

Regler i prompten:
- Bara fakta från profilen och CV:t. Inget hittas på.
- Brevet på engelska om annonsen är på engelska, annars på svenska.
- Kontaktuppgifter läggs in av koden, inte av modellen.
"""

from __future__ import annotations

import io
import json
import re
import sqlite3
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from typing import Callable

from docx import Document
from docx.shared import Pt

from . import llm
from .platsbanken import Annons
from .settings import Profil, Sokinstallningar

Chattfunktion = Callable[..., str]

SV_ORD = {"och", "att", "det", "som", "för", "med", "på", "är", "en", "av", "till", "har", "inte", "vi", "du", "jag"}
EN_ORD = {"and", "the", "to", "of", "in", "for", "with", "is", "we", "you", "our", "are", "that", "will", "a", "this"}


def sprak_for_annons(annons: Annons) -> str:
    """Gissar annonsens språk ('sv' eller 'en') utifrån vanliga ord. Standard: svenska."""
    ord = re.findall(r"[a-zA-ZåäöÅÄÖ]+", f"{annons.rubrik} {annons.beskrivning}".lower())
    sv = sum(1 for w in ord if w in SV_ORD)
    en = sum(1 for w in ord if w in EN_ORD)
    return "en" if en > sv else "sv"


@dataclass
class Dokument:
    annons_id: str
    sprak: str
    profiltext: str = ""
    erfarenhet: list[dict] = field(default_factory=list)  # {titel, arbetsgivare, period, punkter}
    utbildning: list[dict] = field(default_factory=list)  # {examen, larosate, period}
    kompetenser: list[str] = field(default_factory=list)
    sprakkunskaper: list[str] = field(default_factory=list)
    brev: str = ""  # brödtexten; hälsning och avslutning läggs till vid export
    modell: str = ""


SYSTEM_PROMPT = """Du hjälper en person att skriva ett CV och ett personligt brev för ett jobb.

Regler:
- Använd BARA fakta som finns i profilen och CV-utdraget. Hitta aldrig på erfarenheter, examina, siffror eller företag.
- Påstå inget om studier, tillgänglighet, körkort, språk eller tidigare arbete som inte står i profilen. Även om annonsen nämner det.
- Skriv inte att personen uppfyller ett krav i annonsen om det inte står i profilen. Nämn gärna relevant erfarenhet i stället.
- Inga påhittade personlighetsdrag eller omdömen (som "flexibel", "driven", "trivs med") om de inte står i profilen.
- Anpassa texten till annonsen: lyft fram det som matchar kraven, men var ärlig.
- Språk: skriv på det språk som anges i SPRÅK. Svenska ska vara enkel, tydlig och naturlig.
- Brevet: 3 korta stycken, ungefär 180–250 ord. Ingen hälsning och ingen avslutning (de läggs till senare).
- Skriv inga kontaktuppgifter, personnummer eller ålder.

Svara ENDAST med ett JSON-objekt:
{
  "profiltext": "2–3 meningar om personen, anpassade till jobbet",
  "erfarenhet": [{"titel": "", "arbetsgivare": "", "period": "", "punkter": ["..."]}],
  "utbildning": [{"examen": "", "larosate": "", "period": ""}],
  "kompetenser": ["..."],
  "sprakkunskaper": ["..."],
  "brev": "brödtexten i tre stycken, separerade med en tom rad"
}
"""

SPRAK_NAMN = {"sv": "svenska", "en": "engelska"}


def _bygg_prompt(profil: Profil, sok: Sokinstallningar, annons: Annons, sprak: str) -> list[dict]:
    from .matchning import bygg_annonstext, bygg_profiltext  # undviker cirkulär import

    profiltext = bygg_profiltext(profil, sok)
    innehall = (
        f"SPRÅK: {SPRAK_NAMN[sprak]}\n\n"
        f"PROFIL:\n{profiltext}\n\n"
        f"ANNONS:\n{bygg_annonstext(annons)}"
    )
    return [
        {"role": "system", "content": SYSTEM_PROMPT},
        {"role": "user", "content": innehall},
    ]


def skapa_dokument(
    profil: Profil,
    sok: Sokinstallningar,
    annons: Annons,
    modell: str,
    sprak: str | None = None,
    chatt: Chattfunktion = llm.chat,
) -> Dokument:
    sprak = sprak or sprak_for_annons(annons)
    svar = chatt(_bygg_prompt(profil, sok, annons, sprak), modell,
                 json_format=True, temperature=0.4, max_tokens=2000)
    data = llm.parsa_json(svar)

    def lista(namn: str) -> list:
        v = data.get(namn) or []
        return v if isinstance(v, list) else [v]

    return Dokument(
        annons_id=annons.id,
        sprak=sprak,
        profiltext=str(data.get("profiltext", "")).strip(),
        erfarenhet=[e for e in lista("erfarenhet") if isinstance(e, dict)],
        utbildning=[e for e in lista("utbildning") if isinstance(e, dict)],
        kompetenser=[str(k) for k in lista("kompetenser")],
        sprakkunskaper=[str(k) for k in lista("sprakkunskaper")],
        brev=str(data.get("brev", "")).strip(),
        modell=modell,
    )


# --- Lagring -------------------------------------------------------------

def init_tabell(conn: sqlite3.Connection) -> None:
    conn.execute(
        "CREATE TABLE IF NOT EXISTS dokument ("
        " annons_id TEXT PRIMARY KEY,"
        " data TEXT NOT NULL,"
        " uppdaterad TEXT NOT NULL)"
    )
    conn.commit()


def spara(conn: sqlite3.Connection, dok: Dokument) -> None:
    init_tabell(conn)
    conn.execute(
        "INSERT OR REPLACE INTO dokument (annons_id, data, uppdaterad) VALUES (?, ?, ?)",
        (dok.annons_id, json.dumps(asdict(dok), ensure_ascii=False),
         datetime.now().isoformat(timespec="seconds")),
    )
    conn.commit()


def hamta(conn: sqlite3.Connection, annons_id: str) -> Dokument | None:
    init_tabell(conn)
    rad = conn.execute("SELECT data FROM dokument WHERE annons_id = ?", (annons_id,)).fetchone()
    return Dokument(**json.loads(rad["data"])) if rad else None


# --- Word-export ---------------------------------------------------------

ORDLISTA = {
    "sv": {
        "profil": "Profil", "erfarenhet": "Arbetslivserfarenhet", "utbildning": "Utbildning",
        "kompetenser": "Kompetenser", "sprak": "Språk", "referenser": "Referenser",
        "referenser_text": "Lämnas på begäran.", "hej": "Hej,", "halsning": "Med vänliga hälsningar",
        "kontakt": "Kontakt",
    },
    "en": {
        "profil": "Profile", "erfarenhet": "Work experience", "utbildning": "Education",
        "kompetenser": "Skills", "sprak": "Languages", "referenser": "References",
        "referenser_text": "Available on request.", "hej": "Dear hiring team,", "halsning": "Kind regards",
        "kontakt": "Contact",
    },
}


def _kontaktrad(profil: Profil) -> str:
    return " · ".join(x for x in [profil.ort, profil.epost, profil.telefon] if x)


def exportera_cv(profil: Profil, dok: Dokument) -> bytes:
    o = ORDLISTA[dok.sprak]
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.add_heading(profil.namn or "CV", level=0)
    if _kontaktrad(profil):
        doc.add_paragraph(_kontaktrad(profil))
    if dok.profiltext:
        doc.add_heading(o["profil"], level=2)
        doc.add_paragraph(dok.profiltext)
    if dok.erfarenhet:
        doc.add_heading(o["erfarenhet"], level=2)
        for e in dok.erfarenhet:
            titel = " – ".join(x for x in [e.get("titel", ""), e.get("arbetsgivare", "")] if x)
            p = doc.add_paragraph()
            p.add_run(titel).bold = True
            if e.get("period"):
                p.add_run(f"  ({e['period']})")
            for punkt in e.get("punkter") or []:
                doc.add_paragraph(str(punkt), style="List Bullet")
    if dok.utbildning:
        doc.add_heading(o["utbildning"], level=2)
        for u in dok.utbildning:
            rad = " – ".join(x for x in [u.get("examen", ""), u.get("larosate", "")] if x)
            if u.get("period"):
                rad += f"  ({u['period']})"
            doc.add_paragraph(rad, style="List Bullet")
    if dok.kompetenser:
        doc.add_heading(o["kompetenser"], level=2)
        doc.add_paragraph(", ".join(dok.kompetenser))
    if dok.sprakkunskaper or profil.sprak:
        doc.add_heading(o["sprak"], level=2)
        doc.add_paragraph(", ".join(dok.sprakkunskaper or profil.sprak))
    doc.add_heading(o["referenser"], level=2)
    doc.add_paragraph(o["referenser_text"])
    return _till_bytes(doc)


def exportera_brev(profil: Profil, dok: Dokument) -> bytes:
    o = ORDLISTA[dok.sprak]
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_paragraph(profil.namn)
    if _kontaktrad(profil):
        doc.add_paragraph(_kontaktrad(profil))
    doc.add_paragraph(_datum(dok.sprak))
    doc.add_paragraph(o["hej"])
    for stycke in [s.strip() for s in dok.brev.split("\n\n") if s.strip()]:
        doc.add_paragraph(stycke)
    doc.add_paragraph(o["halsning"] + ",")
    doc.add_paragraph(profil.namn)
    return _till_bytes(doc)


def _datum(sprak: str) -> str:
    if sprak == "en":
        return date.today().strftime("%B %d, %Y")
    manader = ["januari", "februari", "mars", "april", "maj", "juni", "juli",
               "augusti", "september", "oktober", "november", "december"]
    d = date.today()
    return f"{d.day} {manader[d.month - 1]} {d.year}"


def _till_bytes(doc: Document) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --- Redigering som text ----------------------------------------------------
# Erfarenhet: ett block per jobb, separerade med tom rad.
#   Titel | Arbetsgivare | Period
#   - punkt
#   - punkt
# Utbildning: en rad per examen: Examen | Lärosäte | Period

def erfarenhet_till_text(erf: list[dict]) -> str:
    block = []
    for e in erf:
        rad = " | ".join([e.get("titel", ""), e.get("arbetsgivare", ""), e.get("period", "")])
        punkter = [f"- {p}" for p in (e.get("punkter") or [])]
        block.append("\n".join([rad, *punkter]))
    return "\n\n".join(block)


def erfarenhet_fran_text(text: str) -> list[dict]:
    resultat = []
    for block in re.split(r"\n\s*\n", text.strip()):
        rader = [r.strip() for r in block.splitlines() if r.strip()]
        if not rader:
            continue
        falt = [f.strip() for f in rader[0].split("|")] + ["", "", ""]
        punkter = [r[1:].strip() for r in rader[1:] if r.startswith("-")]
        resultat.append({"titel": falt[0], "arbetsgivare": falt[1], "period": falt[2], "punkter": punkter})
    return resultat


def utbildning_till_text(utb: list[dict]) -> str:
    return "\n".join(" | ".join([u.get("examen", ""), u.get("larosate", ""), u.get("period", "")]) for u in utb)


def utbildning_fran_text(text: str) -> list[dict]:
    resultat = []
    for rad in text.splitlines():
        if not rad.strip():
            continue
        falt = [f.strip() for f in rad.split("|")] + ["", "", ""]
        resultat.append({"examen": falt[0], "larosate": falt[1], "period": falt[2]})
    return resultat
