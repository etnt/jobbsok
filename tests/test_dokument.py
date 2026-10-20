import io
import json
from pathlib import Path

from docx import Document as LaddaDoc

from jobbsok import dokument
from jobbsok.platsbanken import Annons
from jobbsok.settings import Profil, Sokinstallningar
from jobbsok.storage import connect

FIXTURE = Path(__file__).parent / "fixtures" / "annons.json"


def _annons(rubrik="Lagerarbetare", beskr="Vi söker en person som är noggrann och har god vilja att lära sig."):
    a = Annons.fran_api(json.loads(FIXTURE.read_text(encoding="utf-8")))
    a.rubrik, a.beskrivning = rubrik, beskr
    return a


def _profil():
    return Profil(namn="Anna Andersson", ort="Solna", epost="anna@example.se",
                  telefon="070-111 22 33", sprak=["Svenska (modersmål)", "Engelska (flytande)"],
                  utbildning="Byggingenjör, KTH", kompetenser=["AutoCAD"])


def _svar(sprak_brev="Jag vill söka tjänsten."):
    def chatt(messages, modell, **kw):
        assert "SPRÅK:" in messages[1]["content"]
        return json.dumps({
            "profiltext": "Byggingenjör med intresse för praktiskt arbete.",
            "erfarenhet": [{"titel": "Snickare", "arbetsgivare": "AB X", "period": "2023",
                            "punkter": ["Byggde stommar"]}],
            "utbildning": [{"examen": "Byggingenjör", "larosate": "KTH", "period": "2022–2025"}],
            "kompetenser": ["AutoCAD"], "sprakkunskaper": ["Svenska", "Engelska"],
            "brev": f"Första stycket.\n\n{sprak_brev}\n\nTack för att du läser.",
        })
    return chatt


def test_sprak_detektion():
    assert dokument.sprak_for_annons(_annons()) == "sv"
    en = _annons("Warehouse worker", "We are looking for someone who is reliable and will work with our team.")
    assert dokument.sprak_for_annons(en) == "en"


def test_skapa_dokument_och_prompt_utan_kontaktuppgifter():
    seen = {}

    def chatt(messages, modell, **kw):
        seen["prompt"] = messages[1]["content"]
        return _svar()(messages, modell, **kw)
    dok = dokument.skapa_dokument(_profil(), Sokinstallningar(), _annons(), "m", chatt=chatt)
    assert dok.sprak == "sv" and dok.modell == "m"
    assert "anna@example.se" not in seen["prompt"] and "070-111" not in seen["prompt"]
    assert dok.erfarenhet[0]["titel"] == "Snickare"


def test_spara_och_hamta(tmp_path):
    conn = connect(tmp_path / "t.sqlite3")
    dok = dokument.skapa_dokument(_profil(), Sokinstallningar(), _annons(), "m", chatt=_svar())
    dokument.spara(conn, dok)
    assert dokument.hamta(conn, dok.annons_id) == dok
    assert dokument.hamta(conn, "saknas") is None


def test_cv_docx_innehaller_namn_kontakt_och_punkter():
    dok = dokument.skapa_dokument(_profil(), Sokinstallningar(), _annons(), "m", chatt=_svar())
    doc = LaddaDoc(io.BytesIO(dokument.exportera_cv(_profil(), dok)))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Anna Andersson" in text and "anna@example.se" in text
    assert "Arbetslivserfarenhet" in text and "Byggde stommar" in text
    assert "Lämnas på begäran." in text


def test_brev_docx_har_hals_och_avslutning():
    dok = dokument.skapa_dokument(_profil(), Sokinstallningar(), _annons(), "m", chatt=_svar())
    doc = LaddaDoc(io.BytesIO(dokument.exportera_brev(_profil(), dok)))
    texter = [p.text for p in doc.paragraphs]
    assert "Hej," in texter
    assert "Med vänliga hälsningar," in texter
    assert "Tack för att du läser." in texter


def test_engelskt_brev_pa_engelska():
    dok = dokument.skapa_dokument(_profil(), Sokinstallningar(), _annons(), "m", sprak="en", chatt=_svar())
    doc = LaddaDoc(io.BytesIO(dokument.exportera_brev(_profil(), dok)))
    texter = [p.text for p in doc.paragraphs]
    assert "Dear hiring team," in texter and "Kind regards," in texter


def test_brevstext_delas_i_stycken():
    dok = dokument.skapa_dokument(_profil(), Sokinstallningar(), _annons(), "m", chatt=_svar())
    doc = LaddaDoc(io.BytesIO(dokument.exportera_brev(_profil(), dok)))
    texter = [p.text for p in doc.paragraphs]
    assert "Första stycket." in texter and "Jag vill söka tjänsten." in texter
